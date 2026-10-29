import os
import sys
import time
import docx
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
import docx2pdf
import pypdf

from doc_styles import setup_document
from front_matter import (
    create_cover_page, create_certificate_page,
    create_acknowledgement_page, create_abstract_page,
    create_table_of_contents_page, create_list_of_figures_page,
    create_list_of_tables_page, create_abbreviations_page
)
from ch1_intro import build_chapter_1
from ch2_lit_survey import build_chapter_2
from ch3_methodology import build_chapter_3
from ch4_results import build_chapter_4
from ch5_adv_lim_app import build_chapter_5
from ch6_conclusion import build_chapter_6
from back_matter import build_references, build_appendices

def add_page_number_to_paragraph(paragraph):
    """Insert standard OpenXML PAGE field inside paragraph with Times New Roman 10pt."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFont = OxmlElement('w:rFonts')
    rFont.set(qn('w:ascii'), 'Times New Roman')
    rFont.set(qn('w:hAnsi'), 'Times New Roman')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20') # 10 pt
    rPr.append(rFont)
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fldSimple.append(r)
    paragraph._p.append(fldSimple)

def set_section_page_numbering(section, start=1, num_format='lowerRoman'):
    """Set section page numbering restart and format strictly before w:docGrid in OpenXML sectPr."""
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), str(start))
    pgNumType.set(qn('w:fmt'), num_format)
    docGrid = sectPr.find(qn('w:docGrid'))
    if docGrid is not None:
        docGrid.addprevious(pgNumType)
    else:
        sectPr.append(pgNumType)

def main():
    docx_path = "AI_Student_Agent_Black_Book_Final.docx"
    pdf_path = "AI_Student_Agent_Black_Book_Final.pdf"

    print("=" * 65, flush=True)
    print(" Compiling AI Student Agent Capstone Black Book...", flush=True)
    print("=" * 65, flush=True)

    doc = setup_document()

    # Section 1: Cover Page and Certificate
    print(" -> Building Section 1: Cover Page & Certificate...", flush=True)
    create_cover_page(doc, logo_path="nmims_logo.png")
    create_certificate_page(doc, logo_path="nmims_logo.png")

    # Section 2: Preliminary Pages (Lowercase Roman: i, ii, iii...)
    print(" -> Building Section 2: Preliminary Pages (Lowercase Roman)...", flush=True)
    sec2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False
    set_section_page_numbering(sec2, start=1, num_format='lowerRoman')

    p_f2 = sec2.footer.paragraphs[0]
    p_f2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_to_paragraph(p_f2)

    create_acknowledgement_page(doc)
    create_abstract_page(doc)
    create_table_of_contents_page(doc)
    create_list_of_figures_page(doc)
    create_list_of_tables_page(doc)
    create_abbreviations_page(doc)

    # Section 3: Chapters 1 to 6, References, Appendices (Arabic: 1, 2, 3...)
    print(" -> Building Section 3: Chapters 1-6, References, Appendices...", flush=True)
    sec3 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec3.header.is_linked_to_previous = False
    sec3.footer.is_linked_to_previous = False
    set_section_page_numbering(sec3, start=1, num_format='decimal')

    # Running Header
    p_h3 = sec3.header.paragraphs[0]
    r1 = p_h3.add_run("AI Student Agent: An Autonomous AI Learning Assistant for Real-Time Online Classes\t\t")
    r2 = p_h3.add_run("Academic Year 2026-2027")
    for r in [r1, r2]:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)

    # Footer
    p_f3 = sec3.footer.paragraphs[0]
    p_f3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number_to_paragraph(p_f3)

    print("    * Chapter 1: Introduction", flush=True)
    build_chapter_1(doc)

    print("    * Chapter 2: Literature Survey", flush=True)
    build_chapter_2(doc)

    print("    * Chapter 3: Methodology and Implementation", flush=True)
    build_chapter_3(doc)

    print("    * Chapter 4: Results and Analysis", flush=True)
    build_chapter_4(doc)

    print("    * Chapter 5: Advantages, Limitations and Applications", flush=True)
    build_chapter_5(doc)

    print("    * Chapter 6: Conclusion and Future Scope", flush=True)
    build_chapter_6(doc)

    print("    * References (IEEE Format)", flush=True)
    build_references(doc)

    print("    * Appendices I to IV", flush=True)
    build_appendices(doc)

    # Save DOCX
    t0 = time.time()
    doc.save(docx_path)
    print(f"\n [SUCCESS] Saved '{docx_path}' in {time.time()-t0:.2f}s ({os.path.getsize(docx_path):,} bytes)", flush=True)

    # Export to PDF via docx2pdf
    print(f"\n -> Exporting to PDF via docx2pdf: '{pdf_path}' ...", flush=True)
    if os.path.exists(pdf_path):
        try:
            os.remove(pdf_path)
        except Exception:
            pass

    t_pdf = time.time()
    docx2pdf.convert(docx_path, pdf_path)
    print(f" [SUCCESS] Exported '{pdf_path}' in {time.time()-t_pdf:.2f}s ({os.path.getsize(pdf_path):,} bytes)", flush=True)

    # Verify with pypdf
    reader = pypdf.PdfReader(pdf_path)
    print(f" Total PDF Pages: {len(reader.pages)}", flush=True)
    print("=" * 65, flush=True)

if __name__ == "__main__":
    main()
