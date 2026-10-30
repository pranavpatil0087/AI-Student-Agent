from pptx import Presentation

prs = Presentation(r'C:\Users\ganesh jadhav\Downloads\Sample PPT format External Review.pptx')

print("=== INSPECTING SAMPLE PPT FORMAT EXTERNAL REVIEW ===")
for i, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {i+1} --- (shapes: {len(slide.shapes)})")
    for j, shp in enumerate(slide.shapes):
        pos = f"pos=({shp.left/914400:.2f}, {shp.top/914400:.2f}), size=({shp.width/914400:.2f}, {shp.height/914400:.2f})"
        txt = ""
        font_info = ""
        if shp.has_text_frame:
            txt = shp.text_frame.text.replace('\n', ' | ')[:70]
            if shp.text_frame.paragraphs:
                p = shp.text_frame.paragraphs[0]
                fn = p.font.name or "None"
                fs = p.font.size.pt if p.font.size else "None"
                font_info = f"[{fn}, {fs}pt]"
        print(f"  [{j}] {shp.name} | {shp.shape_type} | {pos} | {font_info} | {txt}")
