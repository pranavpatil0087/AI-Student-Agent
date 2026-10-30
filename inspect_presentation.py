from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation(r'C:\Users\ganesh jadhav\Downloads\MedIntel_AI_Final_Presentation.pptx')

print(f"Slide Width: {prs.slide_width / 914400} in, Slide Height: {prs.slide_height / 914400} in")

slides_to_check = [0, 1, 3, 5, 7, 9, 14, 19, 23, 26]
for s_idx in slides_to_check:
    s = prs.slides[s_idx]
    print(f"\n=================== SLIDE {s_idx+1} ===================")
    print(f"Total shapes: {len(s.shapes)}")
    for i, shp in enumerate(s.shapes):
        t_type = str(shp.shape_type)
        pos = f"pos=({shp.left/914400:.2f}, {shp.top/914400:.2f}), size=({shp.width/914400:.2f}, {shp.height/914400:.2f})"
        txt_prev = ""
        font_info = ""
        if shp.has_text_frame:
            txt_prev = shp.text_frame.text.replace('\n', ' ')[:60]
            if shp.text_frame.paragraphs:
                p = shp.text_frame.paragraphs[0]
                fn = p.font.name or "def"
                fs = p.font.size.pt if p.font.size else "def"
                fc = str(p.font.color.rgb) if p.font.color and hasattr(p.font.color, 'rgb') and p.font.color.rgb else "def"
                font_info = f"[{fn}, {fs}pt, {fc}]"
        elif shp.has_table:
            txt_prev = f"TABLE {len(shp.table.rows)}x{len(shp.table.columns)}"
        print(f"  [{i}] {shp.name} | {t_type} | {pos} | {font_info} | {txt_prev}")
