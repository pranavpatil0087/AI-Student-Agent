from pptx import Presentation

prs = Presentation(r'C:\Users\ganesh jadhav\Downloads\AI_Student_Agent_Final_Presentation.pptx')
print(f"Total slides: {len(prs.slides)}")
for idx, s in enumerate(prs.slides):
    texts = []
    for shp in s.shapes:
        if shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                t = p.text.strip()
                if t and len(t) < 100:
                    texts.append(t)
    print(f"\n--- Slide {idx+1} ---")
    for t in texts[:6]:
        print(f"  {t}")
