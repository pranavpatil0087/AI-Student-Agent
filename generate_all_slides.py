import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Colors matching MedIntel AI
COLOR_RED = RGBColor(188, 49, 44)       # #BC312C (NMIMS Brand Red)
COLOR_BLUE = RGBColor(46, 95, 138)      # #2E5F8A (Steel Navy)
COLOR_GREEN = RGBColor(63, 143, 92)     # #3F8F5C (Forest Green)
COLOR_AMBER = RGBColor(201, 138, 44)    # #C98A2C (Amber Gold)
COLOR_DARK = RGBColor(51, 51, 51)       # #333333 (Charcoal Dark Text)
COLOR_MUTED = RGBColor(102, 102, 102)   # #666666 (Muted Gray Text)
COLOR_CARD_BG = RGBColor(244, 246, 248) # #F4F6F8 (Light Neutral Fill)
COLOR_WHITE = RGBColor(255, 255, 255)   # #FFFFFF
COLOR_BORDER = RGBColor(208, 208, 208)  # #D0D0D0 (Subtle Gray Border)
COLOR_LIGHT_RED = RGBColor(253, 242, 242)

FONT_FAMILY = "Times New Roman"

def create_deck():
    template_path = r'C:\Users\ganesh jadhav\Downloads\MedIntel_AI_Final_Presentation.pptx'
    prs = Presentation(template_path)
    
    # Remove all existing slides to preserve slide master & image1.jpeg background
    for i in range(len(prs.slides) - 1, -1, -1):
        rId = prs.slides._sldIdLst[i].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[i]
        
    blank_layout = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]
    
    # Helper: Add Slide Header & Footer
    def add_header(slide, title_text, slide_num):
        # Title text box
        tb = slide.shapes.add_textbox(Inches(0.50), Inches(1.15), Inches(8.50), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        
        # Slide number
        num_box = slide.shapes.add_textbox(Inches(7.50), Inches(6.92), Inches(2.00), Inches(0.35))
        ntf = num_box.text_frame
        ntf.margin_left = ntf.margin_top = ntf.margin_right = ntf.margin_bottom = 0
        np = ntf.paragraphs[0]
        np.alignment = PP_ALIGN.RIGHT
        np.text = str(slide_num)
        np.font.name = FONT_FAMILY
        np.font.size = Pt(13)
        np.font.color.rgb = COLOR_MUTED

    # Helper: Add Card Box
    def add_card(slide, left, top, width, height, fill_color=COLOR_CARD_BG, border_color=COLOR_BORDER, border_width=Pt(1)):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = border_width
        else:
            card.line.fill.background()
        return card

    # Helper: Add Numbered Badge
    def add_badge(slide, left, top, size, number_str, color=COLOR_RED):
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
        oval.fill.solid()
        oval.fill.fore_color.rgb = color
        oval.line.fill.background()
        tf = oval.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = str(number_str)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_FAMILY
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        return oval

    # Helper: Format Table
    def style_table(table, col_widths=None):
        if col_widths:
            for i, w in enumerate(col_widths):
                table.columns[i].width = w
        for col_idx, cell in enumerate(table.rows[0].cells):
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_RED
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT_FAMILY
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = COLOR_WHITE
                p.alignment = PP_ALIGN.CENTER
        for r_idx in range(1, len(table.rows)):
            row = table.rows[r_idx]
            bg_color = COLOR_WHITE if r_idx % 2 == 1 else COLOR_CARD_BG
            for cell in row.cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg_color
                for p in cell.text_frame.paragraphs:
                    p.font.name = FONT_FAMILY
                    p.font.size = Pt(10)
                    p.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    # NMIMS Logo
    logo_path = 'extracted_media/image1.png'
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(4.50), Inches(0.95), Inches(1.00), Inches(1.04))
    
    # Department Header
    dept_box = s1.shapes.add_textbox(Inches(1.00), Inches(2.05), Inches(8.00), Inches(0.80))
    dtf = dept_box.text_frame
    dtf.word_wrap = True
    dp1 = dtf.paragraphs[0]
    dp1.alignment = PP_ALIGN.CENTER
    dp1.text = "Department of Information Technology"
    dp1.font.name = FONT_FAMILY
    dp1.font.size = Pt(15)
    dp1.font.bold = True
    dp1.font.color.rgb = COLOR_DARK
    dp2 = dtf.add_paragraph()
    dp2.alignment = PP_ALIGN.CENTER
    dp2.text = "Mukesh Patel School of Technology Management & Engineering"
    dp2.font.name = FONT_FAMILY
    dp2.font.size = Pt(13)
    dp2.font.color.rgb = COLOR_MUTED
    dp3 = dtf.add_paragraph()
    dp3.alignment = PP_ALIGN.CENTER
    dp3.text = "SVKM's NMIMS Shirpur Campus | Academic Year 2026-27"
    dp3.font.name = FONT_FAMILY
    dp3.font.size = Pt(12)
    dp3.font.color.rgb = COLOR_MUTED

    # Title Box
    title_box = s1.shapes.add_textbox(Inches(0.60), Inches(3.05), Inches(8.80), Inches(1.60))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    tp1 = ttf.paragraphs[0]
    tp1.alignment = PP_ALIGN.CENTER
    tp1.text = "Final Project Presentation"
    tp1.font.name = FONT_FAMILY
    tp1.font.size = Pt(16)
    tp1.font.color.rgb = COLOR_RED
    tp1.font.bold = True
    tp2 = ttf.add_paragraph()
    tp2.alignment = PP_ALIGN.CENTER
    tp2.text = "on"
    tp2.font.name = FONT_FAMILY
    tp2.font.size = Pt(13)
    tp2.font.color.rgb = COLOR_MUTED
    tp3 = ttf.add_paragraph()
    tp3.alignment = PP_ALIGN.CENTER
    tp3.text = "“AI Student Agent: An Autonomous AI Learning Assistant for Real-Time Online Classes”"
    tp3.font.name = FONT_FAMILY
    tp3.font.size = Pt(21)
    tp3.font.bold = True
    tp3.font.color.rgb = COLOR_DARK

    # Students & Mentor
    info_box = s1.shapes.add_textbox(Inches(0.80), Inches(4.85), Inches(8.40), Inches(1.80))
    itf = info_box.text_frame
    itf.word_wrap = True
    
    ip1 = itf.paragraphs[0]
    ip1.alignment = PP_ALIGN.CENTER
    ip1.text = "Presented By:"
    ip1.font.name = FONT_FAMILY
    ip1.font.size = Pt(13)
    ip1.font.bold = True
    ip1.font.color.rgb = COLOR_DARK
    
    ip2 = itf.add_paragraph()
    ip2.alignment = PP_ALIGN.CENTER
    ip2.text = "1. Aman Choudhary (70012300166)     2. Pranav Patil (70012400176)     3. Mohit Patil (70012400178)"
    ip2.font.name = FONT_FAMILY
    ip2.font.size = Pt(12)
    ip2.font.color.rgb = COLOR_DARK
    
    ip3 = itf.add_paragraph()
    ip3.alignment = PP_ALIGN.CENTER
    ip3.text = "\nUnder the Guidance of:"
    ip3.font.name = FONT_FAMILY
    ip3.font.size = Pt(13)
    ip3.font.bold = True
    ip3.font.color.rgb = COLOR_DARK
    
    ip4 = itf.add_paragraph()
    ip4.alignment = PP_ALIGN.CENTER
    ip4.text = "Dr. Pravin Landge (Faculty Mentor, Department of Information Technology)"
    ip4.font.name = FONT_FAMILY
    ip4.font.size = Pt(12)
    ip4.font.color.rgb = COLOR_RED

    print("Created Slide 1: Title")

    prs.save('test_stage1.pptx')
    print("Stage 1 test saved!")

if __name__ == '__main__':
    create_deck()
