from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation(r'C:\Users\ganesh jadhav\Downloads\MedIntel_AI_Final_Presentation.pptx')

print("=== COLOR & STYLE INSPECTION ===")

def inspect_fill_and_line(shp):
    fill_info = "no_fill"
    line_info = "no_line"
    try:
        if shp.fill.type:
            fill_info = f"fill_type={shp.fill.type}"
            if hasattr(shp.fill, 'fore_color') and hasattr(shp.fill.fore_color, 'rgb'):
                fill_info += f", rgb={shp.fill.fore_color.rgb}"
    except Exception as e:
        fill_info = f"fill_err({e})"
    try:
        if shp.line.color:
            line_info = f"line_color={shp.line.color.rgb if hasattr(shp.line.color, 'rgb') else 'theme'}"
    except Exception as e:
        line_info = "no_line"
    return fill_info, line_info

for i in [1, 3, 5, 14, 19, 23]:
    s = prs.slides[i]
    print(f"\n--- Slide {i+1} Details ---")
    for shp in s.shapes:
        f_info, l_info = inspect_fill_and_line(shp)
        if shp.has_table:
            cell = shp.table.cell(0, 0)
            c_fill, _ = inspect_fill_and_line(cell)
            p = cell.text_frame.paragraphs[0]
            print(f"Table Header Cell (0,0): {c_fill}, Text: '{p.text}', Font: {p.font.name}, {p.font.size.pt if p.font.size else 'None'}pt, Bold: {p.font.bold}")
        elif shp.has_text_frame and len(shp.text_frame.text.strip()) > 0:
            p = shp.text_frame.paragraphs[0]
            f_rgb = p.font.color.rgb if p.font and p.font.color and hasattr(p.font.color, 'rgb') else 'None'
            print(f"Shape '{shp.name}' ({shp.shape_type}): {f_info} | {l_info} | Font: {p.font.name}, {p.font.size.pt if p.font.size else 'None'}pt, bold={p.font.bold}, rgb={f_rgb} | Text: '{p.text[:40]}'")
        elif shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            print(f"Card Shape '{shp.name}': {f_info} | {l_info}")
