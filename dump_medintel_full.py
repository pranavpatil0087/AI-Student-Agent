from pptx import Presentation

prs = Presentation(r'C:\Users\ganesh jadhav\Downloads\MedIntel_AI_Final_Presentation.pptx')

with open("medintel_full_text.txt", "w", encoding="utf-8") as out:
    for i, slide in enumerate(prs.slides):
        out.write(f"\n{'='*30} SLIDE {i+1} {'='*30}\n")
        for shp in slide.shapes:
            if shp.has_text_frame:
                txt = shp.text_frame.text.strip()
                if txt:
                    out.write(f"[{shp.name}]:\n{txt}\n\n")
            elif shp.has_table:
                out.write(f"[{shp.name} TABLE {len(shp.table.rows)}x{len(shp.table.columns)}]:\n")
                for r in shp.table.rows:
                    row_txt = " | ".join(c.text_frame.text.strip().replace('\n', ' ') for c in r.cells)
                    out.write(f"  {row_txt}\n")
                out.write("\n")
            elif shp.shape_type == 13:
                out.write(f"[{shp.name} PICTURE]: pos=({shp.left/914400:.2f}, {shp.top/914400:.2f}), size=({shp.width/914400:.2f}, {shp.height/914400:.2f})\n\n")

print("Exported medintel_full_text.txt successfully!")
