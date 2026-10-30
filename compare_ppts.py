from pptx import Presentation

med = Presentation(r'C:\Users\ganesh jadhav\Downloads\MedIntel_AI_Final_Presentation.pptx')
curr = Presentation(r'C:\Users\ganesh jadhav\Downloads\AI_Student_Agent_Final_Presentation.pptx')

print("MEDINTEL SLIDES (Total:", len(med.slides), "):")
for i, s in enumerate(med.slides):
    title = ""
    for shp in s.shapes:
        if shp.name.startswith("Title") or (shp.has_text_frame and "Presentation" in shp.text_frame.text) or (shp.has_text_frame and "Thank You" in shp.text_frame.text):
            title = shp.text_frame.text.split('\n')[0].strip()
            break
        elif shp.has_text_frame and shp.text_frame.text.strip():
            first_line = shp.text_frame.text.split('\n')[0].strip()
            if len(first_line) < 40 and not first_line.isdigit():
                title = first_line
                break
    pic_count = sum(1 for shp in s.shapes if shp.shape_type == 13)
    tbl_count = sum(1 for shp in s.shapes if shp.has_table)
    card_count = sum(1 for shp in s.shapes if shp.shape_type == 1)
    print(f"Slide {i+1}: '{title}' | shapes={len(s.shapes)}, pics={pic_count}, tbls={tbl_count}, cards={card_count}")

print("\nCURRENT AI STUDENT AGENT SLIDES (Total:", len(curr.slides), "):")
for i, s in enumerate(curr.slides):
    title = ""
    for shp in s.shapes:
        if shp.name.startswith("Title") or (shp.has_text_frame and "Presentation" in shp.text_frame.text) or (shp.has_text_frame and "Thank You" in shp.text_frame.text):
            title = shp.text_frame.text.split('\n')[0].strip()
            break
        elif shp.has_text_frame and shp.text_frame.text.strip():
            first_line = shp.text_frame.text.split('\n')[0].strip()
            if len(first_line) < 40 and not first_line.isdigit():
                title = first_line
                break
    pic_count = sum(1 for shp in s.shapes if shp.shape_type == 13)
    tbl_count = sum(1 for shp in s.shapes if shp.has_table)
    card_count = sum(1 for shp in s.shapes if shp.shape_type == 1)
    print(f"Slide {i+1}: '{title}' | shapes={len(s.shapes)}, pics={pic_count}, tbls={tbl_count}, cards={card_count}")
