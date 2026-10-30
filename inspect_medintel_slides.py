import json
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation(r'C:\Users\ganesh jadhav\Downloads\MedIntel_AI_Final_Presentation.pptx')

slides_data = []

for i, slide in enumerate(prs.slides):
    s_info = {
        "slide_index": i + 1,
        "shapes": []
    }
    for j, shp in enumerate(slide.shapes):
        sh_data = {
            "name": shp.name,
            "type": str(shp.shape_type),
            "left": round(shp.left / 914400, 2),
            "top": round(shp.top / 914400, 2),
            "width": round(shp.width / 914400, 2),
            "height": round(shp.height / 914400, 2),
        }
        if shp.has_text_frame:
            paragraphs = []
            for p in shp.text_frame.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    runs = []
                    for r in p.runs:
                        r_col = None
                        if r.font.color and hasattr(r.font.color, 'rgb'):
                            try:
                                r_col = str(r.font.color.rgb)
                            except:
                                pass
                        runs.append({
                            "text": r.text,
                            "bold": r.font.bold,
                            "size": round(r.font.size.pt, 1) if r.font.size else None,
                            "color": r_col,
                            "font": r.font.name
                        })
                    paragraphs.append({
                        "text": p_text[:100],
                        "runs": runs[:3]
                    })
            sh_data["paragraphs"] = paragraphs[:5]
        elif shp.has_table:
            sh_data["table"] = {
                "rows": len(shp.table.rows),
                "cols": len(shp.table.columns)
            }
        s_info["shapes"].append(sh_data)
    slides_data.append(s_info)

with open("medintel_blueprint.json", "w", encoding="utf-8") as f:
    json.dump(slides_data, f, indent=2)

print("Dumped medintel_blueprint.json successfully!")
