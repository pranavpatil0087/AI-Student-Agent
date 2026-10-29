import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from doc_styles import (
    add_body_p, add_chapter_heading, add_section_heading,
    add_subsection_heading, add_table_title, format_table
)

def build_references(doc):
    """Generate the References section formatted strictly according to IEEE standards (2 pages)."""
    add_chapter_heading(doc, "References")
    
    references = [
        "[1] A. Radford, J. W. Kim, T. Xu, G. Brockman, C. McLeavey, and I. Sutskever, \"Robust speech recognition via large-scale weak supervision,\" in Proceedings of the 40th International Conference on Machine Learning (ICML), PMLR, vol. 202, 2023, pp. 28492–28518.",
        "[2] D. Povey, A. Ghoshal, G. Boulianne, L. Burget, O. Glembek, N. Goel, M. Hannemann, P. Motlíček, Y. Qian, P. Schwarz, J. Silovský, G. Stemmer, and K. Veselý, \"The Kaldi speech recognition toolkit,\" in Proceedings of the IEEE Workshop on Automatic Speech Recognition and Understanding (ASRU), 2011, pp. 1–4.",
        "[3] Alpha Cephei, \"Vosk: Offline open-source speech recognition toolkit,\" Alpha Cephei Inc., [Online]. Available: https://alphacephei.com/vosk/. Accessed: Sep. 2026.",
        "[4] Y. Wang, C. S. Leow, A. Kobayashi, T. Utsuro, and H. Nishizaki, \"ExKaldi-RT: A real-time automatic speech recognition extension toolkit of Kaldi,\" arXiv preprint arXiv:2104.01384, 2021.",
        "[5] A. Dewan, M. Ziemski, H. Meylan, L. Concina, and B. Pouliquen, \"Developing automatic verbatim transcripts for international multilingual meetings: An end-to-end solution,\" arXiv preprint arXiv:2309.15609, 2023.",
        "[6] M. T. R. Laskar, X.-Y. Fu, C. Chen, and S. Bhushan TN, \"Building real-world meeting summarization systems using large language models: A practical perspective,\" arXiv preprint arXiv:2310.19233, 2023.",
        "[7] V. Rennard, G. Shang, J. Hunter, and M. Vazirgiannis, \"Abstractive meeting summarization: A survey,\" arXiv preprint arXiv:2208.04163, 2022.",
        "[8] M. Zhong, D. Yin, T. Yu, A. Zaidi, M. Mutuma, R. Jha, A. H. Awadallah, A. Celikyilmaz, Y. Liu, X. Qiu, and D. Radev, \"QMSum: A new benchmark for query-based multi-domain meeting summarization,\" in Proceedings of the 2021 Conference of the North American Chapter of the Association for Computational Linguistics (NAACL), 2021, pp. 5905–5921.",
        "[9] A. Madhavi, A. Chilakamarri, C. Jupudi, S. Madanaboina, and S. Sriram, \"Automatic running notes generation from audio lecture: A systematic review,\" in Proceedings of the International Conference on Machine Intelligence and Smart Innovation, Lecture Notes in Networks and Systems, Springer, Singapore, 2025, pp. 112–128.",
        "[10] K. Taneja, P. Maiti, S. Kakar, P. Guruprasad, S. Rao, and A. K. Goel, \"Jill Watson: A virtual teaching assistant powered by ChatGPT,\" arXiv preprint arXiv:2405.11070, 2024.",
        "[11] K. Wang, J. Ramos, and R. Lawrence, \"ChatEd: A chatbot leveraging ChatGPT for an enhanced learning experience in higher education,\" arXiv preprint arXiv:2401.00052, 2024.",
        "[12] L. Labadze, M. Grigolia, and L. Machaidze, \"Role of AI chatbots in education: Systematic literature review,\" International Journal of Educational Technology in Higher Education, vol. 20, no. 56, pp. 1–19, 2023.",
        "[13] S. Wollny, J. Schneider, D. Di Mitri, J. Weidlich, M. Rittberger, and H. Drachsler, \"Are we there yet? — A systematic literature review on chatbots in education,\" Frontiers in Artificial Intelligence, vol. 4, art. 654924, pp. 1–18, 2021.",
        "[14] V. Bhardwaj and M. Kumar, \"Transforming higher education with robotic process automation: Enhancing efficiency, innovation, and student-centered learning,\" Discover Sustainability, vol. 6, no. 1, pp. 1–21, 2025.",
        "[15] M.-V. Toma, \"Robotic process automation in higher education: A case study on university daily sheets,\" Journal of Applied Computer Science & Mathematics, vol. 17, no. 2, pp. 21–26, 2023.",
        "[16] W. Kuang and A. Baul, \"A real-time attendance system using deep learning face recognition,\" in Proceedings of the ASEE Annual Conference & Exposition, Long Beach, CA, 2021, pp. 1–12.",
        "[17] K.-T. Tran, D. Dao, M.-D. Nguyen, Q.-V. Pham, B. O'Sullivan, and H. D. Nguyen, \"Multi-agent collaboration mechanisms: A survey of LLMs,\" arXiv preprint arXiv:2501.06322, 2025.",
        "[18] OpenAI, \"GPT-4o mini: Advancing cost-efficient intelligence,\" OpenAI Technical Report, 2024. [Online]. Available: https://openai.com/index/gpt-4o-mini-advancing-cost-efficient-intelligence/. Accessed: Oct. 2026.",
    ]

    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.10
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.0)

    doc.add_page_break()

def build_appendices(doc):
    """Generate Appendices A through D strictly fitting on 1 page each (= 4 pages total)."""
    
    # ---------------- APPENDIX A (1 page) ----------------
    add_chapter_heading(doc, "Appendix A: Sample Code Excerpts")
    
    add_body_p(doc, "This appendix compiles representative production source code excerpts illustrating core algorithms implemented in AI Student Agent. All external API credentials have been sanitized.")

    add_subsection_heading(doc, "Excerpt 1: Audio Capture & DSP Bandpass Filtering (audio_processing.py)")
    code_audio = (
        "def capture_audio_from_device():\n"
        "    p = pyaudio.PyAudio()\n"
        "    dev_idx = find_best_audio_device(p, target_name='cable output')\n"
        "    stream = p.open(format=pyaudio.paInt16, channels=1, rate=16000, input=True,\n"
        "                    input_device_index=dev_idx, frames_per_buffer=2000)\n"
        "    b, a = butter(4, [300.0/8000, 3400.0/8000], btype='band')\n"
        "    while is_running:\n"
        "        data = stream.read(2000, exception_on_overflow=False)\n"
        "        master_recording_frames.append(data)\n"
        "        filtered = lfilter(b, a, np.frombuffer(data, dtype=np.int16)).astype(np.int16)\n"
        "        with buffer_lock: audio_buffer += filtered.tobytes()"
    )
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(0)
    p_c1.paragraph_format.space_after = Pt(3)
    p_c1.paragraph_format.line_spacing = 1.0
    p_c1.paragraph_format.left_indent = Inches(0.2)
    r_c1 = p_c1.add_run(code_audio)
    r_c1.font.name = 'Consolas'
    r_c1.font.size = Pt(7.0)

    add_subsection_heading(doc, "Excerpt 2: VOSK Offline Decoding & Sliding Deque (main.py)")
    code_vosk = (
        "def speech_recognition_worker():\n"
        "    recognizer = KaldiRecognizer(Model('model'), 16000)\n"
        "    rolling_history = deque(maxlen=50)\n"
        "    while is_running:\n"
        "        with buffer_lock:\n"
        "            if len(audio_buffer) >= 96000:\n"
        "                chunk = audio_buffer[:96000]; audio_buffer = audio_buffer[96000:]\n"
        "        if chunk and recognizer.AcceptWaveform(chunk):\n"
        "            txt = json.loads(recognizer.Result()).get('text', '').strip()\n"
        "            if txt:\n"
        "                words = txt.split(); rolling_history.extend(words)\n"
        "                evaluate_triggers(words, list(rolling_history)[-15:])"
    )
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.space_before = Pt(0)
    p_c2.paragraph_format.space_after = Pt(3)
    p_c2.paragraph_format.line_spacing = 1.0
    p_c2.paragraph_format.left_indent = Inches(0.2)
    r_c2 = p_c2.add_run(code_vosk)
    r_c2.font.name = 'Consolas'
    r_c2.font.size = Pt(7.0)

    add_subsection_heading(doc, "Excerpt 3: Selenium Google Meet Chat Injection (main.py)")
    code_sel = (
        "def send_message_to_chat(driver, message):\n"
        "    if not driver.find_elements(By.XPATH, \"//div[@aria-label='Chat with everyone']\"):\n"
        "        WebDriverWait(driver, 5).until(EC.element_to_be_clickable(\n"
        "            (By.XPATH, \"//button[contains(@aria-label, 'chat')]\"))).click()\n"
        "    box = WebDriverWait(driver, 5).until(EC.presence_of_element_located(\n"
        "        (By.XPATH, \"//textarea[@aria-label='Send a message to everyone']\")))\n"
        "    box.send_keys(message); box.send_keys(Keys.RETURN)"
    )
    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.space_before = Pt(0)
    p_c3.paragraph_format.space_after = Pt(3)
    p_c3.paragraph_format.line_spacing = 1.0
    p_c3.paragraph_format.left_indent = Inches(0.2)
    r_c3 = p_c3.add_run(code_sel)
    r_c3.font.name = 'Consolas'
    r_c3.font.size = Pt(7.0)

    add_subsection_heading(doc, "Excerpt 4: Notion API Block Publishing (notion_integration.py)")
    code_notion = (
        "def notes_to_notion(api_token, page_id, markdown_text, title='Class Notes'):\n"
        "    notion = Client(auth=api_token); blocks = []\n"
        "    for line in markdown_text.splitlines():\n"
        "        line = line.strip()\n"
        "        if not line: continue\n"
        "        b_type, text = ('heading_1', line[2:]) if line.startswith('# ') else ('paragraph', line)\n"
        "        blocks.append({'object': 'block', 'type': b_type, b_type: {'rich_text': [{'type': 'text', 'text': {'content': text}}]}})\n"
        "    return notion.pages.create(parent={'page_id': page_id}, properties={'title': [{'text': {'content': title}}]}, children=blocks)['id']"
    )
    p_c4 = doc.add_paragraph()
    p_c4.paragraph_format.space_before = Pt(0)
    p_c4.paragraph_format.space_after = Pt(3)
    p_c4.paragraph_format.line_spacing = 1.0
    p_c4.paragraph_format.left_indent = Inches(0.2)
    r_c4 = p_c4.add_run(code_notion)
    r_c4.font.name = 'Consolas'
    r_c4.font.size = Pt(7.0)

    doc.add_page_break()

    # ---------------- APPENDIX B (1 page) ----------------
    add_chapter_heading(doc, "Appendix B: Technical Specifications and Component Data Sheets")
    
    add_body_p(doc, "This appendix compiles technical data sheets and operational specifications for primary software frameworks integrated within AI Student Agent.")

    add_table_title(doc, "Table B.1: VOSK Kaldi Acoustic Model Specifications")
    vosk_specs = [
        ("Parameter / Attribute", "Technical Specification", "Operational Description"),
        ("Model Identifier", "vosk-model-small-en-us-0.15", "Compact American English acoustic and language model"),
        ("Model Size on Disk", "40.2 MB (Compressed) / 112 MB (Extracted)", "Optimized for minimal RAM allocation on consumer PCs"),
        ("Sampling Frequency", "16,000 Hz (16 kHz)", "Standard wideband speech acoustic sampling"),
        ("Acoustic Feature Extractor", "40-dimensional MFCCs + Pitch Features", "Frame-level Mel-frequency cepstral coefficients"),
        ("Acoustic Model Topology", "Factorized TDNN (Time-Delay Neural Net)", "Sub-sampled multi-layer neural network architecture"),
        ("Decoding Framework", "Kaldi Online2 WFST Dynamic Decoder", "Weighted Finite-State Transducer search graph"),
        ("Vocabulary Capacity", "~128,000 Lexical Words", "General-domain English vocabulary coverage"),
        ("Average Decoding Latency", "350 ms – 480 ms per 3.0s Audio Buffer", "Real-time factor (RTF) < 0.25 on Intel Core i5 CPU"),
        ("Operational Cost", "$0.00 (Completely Free & Open Source)", "Runs 100% locally with zero cloud API reliance"),
    ]
    tbl_b1 = doc.add_table(rows=len(vosk_specs), cols=3)
    tbl_b1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b1.autofit = False
    
    b1_widths = [Inches(1.50), Inches(1.80), Inches(2.47)]
    for row in tbl_b1.rows:
        for idx, w in enumerate(b1_widths):
            row.cells[idx].width = w

    for i, row in enumerate(vosk_specs):
        for c, text in enumerate(row):
            tbl_b1.cell(i, c).text = text

    format_table(tbl_b1, header_bg="E8EEF5", font_size=7.5, header_font_size=8.0)

    add_body_p(doc, "")

    add_table_title(doc, "Table B.2: OpenAI GPT-4o-mini Operational Parameter Specifications")
    gpt_specs = [
        ("Parameter / Attribute", "Technical Specification", "Operational Description"),
        ("Model Identifier", "gpt-4o-mini-2024-07-18", "Cost-efficient multimodal frontier reasoning model"),
        ("Context Window", "128,000 Tokens (~96,000 Words)", "Accommodates extended multi-hour lecture transcripts"),
        ("Maximum Output Tokens", "16,384 Tokens", "Enables comprehensive hierarchical study guide synthesis"),
        ("Average Generation Latency", "1.85 s – 2.65 s (Time to Complete)", "Sub-3 second response time suitable for in-class chat"),
        ("Temperature Setting (Q&A)", "0.3 (Low Deterministic Sampling)", "Constrains generation to strictly factual academic answers"),
        ("Temperature Setting (Notes)", "0.4 (Structured Abstractive Synthesis)", "Ensures coherent, pedagogically sound note summaries"),
        ("Input Token Pricing", "$0.150 per 1,000,000 Tokens", "Reduces per-lecture Q&A cost to under $0.001"),
        ("Output Token Pricing", "$0.600 per 1,000,000 Tokens", "Reduces per-lecture note generation cost to under $0.005"),
    ]
    tbl_b2 = doc.add_table(rows=len(gpt_specs), cols=3)
    tbl_b2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b2.autofit = False
    
    b2_widths = [Inches(1.50), Inches(1.80), Inches(2.47)]
    for row in tbl_b2.rows:
        for idx, w in enumerate(b2_widths):
            row.cells[idx].width = w

    for i, row in enumerate(gpt_specs):
        for c, text in enumerate(row):
            tbl_b2.cell(i, c).text = text

    format_table(tbl_b2, header_bg="E8EEF5", font_size=7.5, header_font_size=8.0)

    add_body_p(doc, "")
    doc.add_page_break()

    # ---------------- APPENDIX C (1 page) ----------------
    add_chapter_heading(doc, "Appendix C: Comprehensive List of Components")
    
    add_body_p(doc, "This appendix enumerates the complete inventory of hardware, system drivers, and software libraries utilized across the development of AI Student Agent.")

    add_table_title(doc, "Table C.1: Comprehensive Hardware Component Manifest")
    hw_list = [
        ("Sr.", "Subsystem / Component", "Manufacturer / Model", "Technical Description", "Verified Specification"),
        ("1", "Host System Processor", "Intel Corporation", "Quad-Core x86-64 with AVX2", "Core i5-1035G1 @ 1.00GHz up to 3.60GHz"),
        ("2", "System Memory (RAM)", "Crucial / Kingston", "DDR4 2666MHz Synchronous RAM", "8.0 GB / 16.0 GB DDR4 Memory"),
        ("3", "System Mass Storage", "Western Digital / Samsung", "NVMe PCIe Gen 3.0 x4 SSD", "256 GB / 512 GB M.2 NVMe SSD"),
        ("4", "Audio Rendering Chipset", "Realtek Semiconductor", "HD Audio Codec (ALC3204)", "24-bit, 192 kHz Hardware DAC/ADC"),
        ("5", "Virtual Audio Cable Driver", "VB-Audio Software", "Virtual Audio Loopback Bridge", "VB-Audio Cable Driver v1.0.3.8"),
        ("6", "Network Controller", "Intel / Realtek", "Wi-Fi 6 (802.11ax) / Gigabit Ethernet", "100 Mbps Broadband Connection"),
        ("7", "Display Subsystem", "Intel UHD Graphics", "Integrated System Video Graphics", "Full HD 1920x1080 Resolution Display"),
    ]
    tbl_c1 = doc.add_table(rows=len(hw_list), cols=5)
    tbl_c1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_c1.autofit = False
    
    c1_widths = [Inches(0.45), Inches(1.45), Inches(1.30), Inches(1.50), Inches(1.07)]
    for row in tbl_c1.rows:
        for idx, w in enumerate(c1_widths):
            row.cells[idx].width = w

    for i, row in enumerate(hw_list):
        for c, text in enumerate(row):
            tbl_c1.cell(i, c).text = text

    format_table(tbl_c1, header_bg="E8EEF5", font_size=7.0, header_font_size=7.5)

    add_body_p(doc, "")

    add_table_title(doc, "Table C.2: Python Package Ecosystem and Installed Dependency Manifest")
    pkg_list = [
        ("Sr.", "Package Identifier", "Installed Version", "Primary Functional Dependency", "License"),
        ("1", "python", "3.13.14 / 3.10+", "Core runtime programming language environment", "PSF License"),
        ("2", "selenium", "4.49.0", "Browser automation, Chrome navigation, DOM chat injection", "Apache 2.0"),
        ("3", "vosk", "0.3.45", "Offline Kaldi speech-to-text decoding on CPU", "Apache 2.0"),
        ("4", "PyAudio", "0.2.14", "Real-time portaudio bindings for virtual cable capture", "MIT License"),
        ("5", "sounddevice", "0.5.5", "Alternative host audio querying and PCM streaming", "MIT License"),
        ("6", "scipy", "1.17.0", "Butterworth digital bandpass filtering and DSP routines", "BSD-3-Clause"),
        ("7", "numpy", "2.3.5", "Vector operations on PCM audio byte buffers", "BSD-3-Clause"),
        ("8", "Flask", "3.1.3", "Internal HTTP REST server for cross-thread event broker", "BSD-3-Clause"),
        ("9", "openai", "3.14.0", "SDK for GPT-4o-mini reasoning and Whisper STT", "Apache 2.0"),
        ("10", "google-genai", "2.23.0", "Official Google GenAI SDK for Gemini fallback", "Apache 2.0"),
        ("11", "notion-client", "3.1.0", "SDK for Notion REST API page and block creation", "MIT License"),
    ]
    tbl_c2 = doc.add_table(rows=len(pkg_list), cols=5)
    tbl_c2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_c2.autofit = False
    
    c2_widths = [Inches(0.45), Inches(1.45), Inches(1.00), Inches(2.00), Inches(0.87)]
    for row in tbl_c2.rows:
        for idx, w in enumerate(c2_widths):
            row.cells[idx].width = w

    for i, row in enumerate(pkg_list):
        for c, text in enumerate(row):
            tbl_c2.cell(i, c).text = text

    format_table(tbl_c2, header_bg="E8EEF5", font_size=7.0, header_font_size=7.5)

    add_body_p(doc, "")
    doc.add_page_break()

    # ---------------- APPENDIX D (1 page) ----------------
    add_chapter_heading(doc, "Appendix D: List of Papers Presented and Published")
    
    add_body_p(doc, "In accordance with academic guidelines of the Information Technology Department, MPSTME, SVKM’s NMIMS, this appendix documents the scholarly manuscript prepared from the capstone project work.")

    add_subsection_heading(doc, "1. Research Paper Prepared for Publication")
    
    add_body_p(doc, "Title: “AI Student Agent: An Autonomous Multi-Threaded Architecture for Real-Time Online Classroom Participation and Lecture Note Synthesis”", bold_prefix="• ")
    add_body_p(doc, "Authors: Aman Choudhary, Pranav Patil, Mohit Patil, and Dr. Pravin Landge", bold_prefix="• ")
    add_body_p(doc, "Institutional Affiliation: Department of Information Technology, MPSTME, SVKM’s NMIMS, Shirpur Campus, Maharashtra, India", bold_prefix="• ")
    add_body_p(doc, "Target Venue: IEEE International Conference on Advanced Learning Technologies (ICALT) / Springer MISI (2027)", bold_prefix="• ")
    add_body_p(doc, "Current Status: Manuscript finalized and under institutional review prior to formal peer-review submission.", bold_prefix="• ")

    add_subsection_heading(doc, "2. Executive Abstract of the Prepared Manuscript")
    p_abs = doc.add_paragraph()
    p_abs.paragraph_format.line_spacing = 1.15
    p_abs.paragraph_format.space_before = Pt(2)
    p_abs.paragraph_format.space_after = Pt(2)
    p_abs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_abs = p_abs.add_run("“Synchronous online education conducted over digital conferencing platforms imposes acute cognitive multitasking demands on engineering students, who must concurrently comprehend complex instructional concepts, record structured notes, and monitor class communication channels for instructor cold-calls. This paper presents AI Student Agent, an autonomous, multi-threaded AI assistant engineered for Google Meet virtual classrooms. The system combines local offline speech-to-text decoding (VOSK/Kaldi) operating via virtual audio loopback routing (VB-Audio Virtual Cable) with cloud generative language models (OpenAI GPT-4o-mini and Whisper). By maintaining an in-memory sliding discourse deque, the agent evaluates acoustic confidence and normalized Levenshtein similarity to detect student name callouts and spoken technical inquiries. Upon trigger verification, the agent synthesizes contextually grounded answers and submits them to the meeting chat via Selenium browser automation within 3.5 seconds. Following session completion, lossless 16 kHz PCM audio is transcribed via Whisper and converted into hierarchical lecture summaries synchronized directly to Notion via REST APIs. Empirical validation across live university lectures demonstrates sub-500 ms offline STT latency, zero continuous cloud streaming costs, a constrained 1.25 GB RAM footprint, and over 98% operational compute savings over pure cloud streaming architectures.”")
    r_abs.font.name = 'Times New Roman'
    r_abs.font.size = Pt(8.5)
    r_abs.font.italic = True
